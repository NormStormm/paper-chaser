"""Step 3: write the results to an Excel workbook that an AP clerk can work from."""
from __future__ import annotations

import pathlib
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .checks import Decision, coi_problems

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HOLD_FILL = PatternFill("solid", fgColor="FDE2E1")
NOTE_FILL = PatternFill("solid", fgColor="FFF4CE")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
MONEY = '"$"#,##0.00'


def _d(s):
    """Real Excel dates (sortable, filterable) instead of text; leaves anything unparseable as it was."""
    try:
        return date.fromisoformat(s) if s else None
    except (TypeError, ValueError):
        return s


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEADER_FILL, Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def _style_row(ws, fill=None, money_cols=(), wrap_cols=(), date_cols=()):
    r = ws.max_row
    for c in ws[r]:
        if fill:
            c.fill = fill
        c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
        if c.column in money_cols:
            c.number_format = MONEY
        if c.column in date_cols:
            c.number_format = "yyyy-mm-dd"


def write_workbook(path: pathlib.Path, decisions: list[Decision], extractions: list[dict], as_of: date, model: str) -> dict:
    wb = Workbook()
    wb.remove(wb.active)
    invoices = [d for d in decisions if d.doc_type == "INVOICE"]
    approved = [d for d in invoices if d.decision == "APPROVED"]
    held = [d for d in decisions if d.decision == "HOLD" and d.doc_type != "COI"]
    cois = [d for d in decisions if d.doc_type == "COI"]
    total = lambda ds: round(sum(d.data.get("total") or 0 for d in ds), 2)

    # Summary
    ws = _sheet(wb, "Summary", ["", ""], [42, 22])
    ws.cell(row=1, column=1, value="Paper Chaser run")
    rows = [("Run on", datetime.now().strftime("%Y-%m-%d %H:%M")), ("Checked as of", as_of.isoformat()), ("Model", model),
            ("Documents read", len(decisions)), ("Invoices approved to enter", len(approved)),
            ("  Amount approved", total(approved)), ("On hold for review", len(held)),
            ("  Invoice amount on hold", total([d for d in held if d.doc_type == "INVOICE"])),
            ("Insurance certificates recorded", len(cois)),
            ("Approved with a note to read", sum(1 for d in approved if d.notes))]
    for label, val in rows:
        ws.append([label, val])
        if label.strip().startswith("Amount") or label.strip().startswith("Invoice amount"):
            ws.cell(row=ws.max_row, column=2).number_format = MONEY
    ws.append([])
    ws.append(["Every amount and status comes from plain-code checks. Claude only reads the documents."])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="666666")

    # Approved
    ws = _sheet(wb, "Approved", ["File", "Vendor", "Invoice #", "Invoice date", "Due date", "PO #", "Total", "Note"],
                [30, 30, 15, 13, 13, 12, 13, 60])
    for d in approved:
        x = d.data
        ws.append([d.file, d.vendor, x.get("invoice_number"), _d(x.get("invoice_date")), _d(x.get("due_date")), x.get("po_number"),
                   x.get("total"), "; ".join(t for _, t in d.notes)])
        _style_row(ws, NOTE_FILL if d.notes else None, money_cols=(7,), wrap_cols=(8,), date_cols=(4, 5))

    # Exceptions
    ws = _sheet(wb, "Exceptions", ["File", "Type", "Vendor", "Invoice #", "PO #", "Total", "Why it's on hold", "Reason codes"],
                [30, 10, 30, 15, 12, 13, 70, 26])
    for d in held:
        x = d.data
        ws.append([d.file, d.doc_type, d.vendor, x.get("invoice_number"), x.get("po_number"), x.get("total"),
                   "\n".join(t for _, t in d.reasons), ", ".join(c for c, _ in d.reasons)])
        _style_row(ws, HOLD_FILL, money_cols=(6,), wrap_cols=(7,))

    # Insurance
    ws = _sheet(wb, "Insurance", ["File", "Insured vendor", "Insurer", "Policy #", "Expires", "Per occurrence", "Status"],
                [36, 30, 30, 16, 12, 15, 60])
    for d in cois:
        x = d.data
        holds, notes = coi_problems(x, as_of)
        status = "; ".join(t for _, t in d.notes) or "OK"
        ws.append([d.file, d.vendor, x.get("insurer"), x.get("policy_number"), _d(x.get("policy_expiration")),
                   x.get("each_occurrence_limit"), status])
        _style_row(ws, HOLD_FILL if holds else NOTE_FILL if notes else OK_FILL, money_cols=(6,), wrap_cols=(7,), date_cols=(5,))

    # Audit log: one row per document, what was read and what was decided.
    ws = _sheet(wb, "Audit Log", ["Time", "File", "Read as", "Decision", "Details", "Model", "Tokens in/out", "Seconds", "Reused saved read?"],
                [17, 30, 10, 10, 70, 18, 13, 9, 11])
    by_file = {e["file"]: e for e in extractions}
    for d in decisions:
        e = by_file.get(d.file, {})
        details = "; ".join(t for _, t in d.reasons + d.notes) or ("Recorded" if d.doc_type == "COI" else "All checks passed")
        if (d.data or {}).get("unclear"):
            details += f"  [Claude: {d.data['unclear']}]"
        ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), d.file, d.doc_type, "RECORDED" if d.doc_type == "COI" and d.decision != "HOLD" else d.decision,
                   details, e.get("model"), f"{e.get('input_tokens', 0)}/{e.get('output_tokens', 0)}", e.get("seconds"),
                   "yes" if e.get("cached") else "no"])
        _style_row(ws, wrap_cols=(5,))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"approved": len(approved), "held": len(held), "cois": len(cois), "amount_approved": total(approved)}
