"""Turns an output workbook into a Markdown page, so its contents can be read on GitHub (which can't preview .xlsx).

    .venv/bin/python tools/preview_workbook.py docs/sample-output.xlsx docs/sample-output.md
"""
import datetime
import sys

from openpyxl import load_workbook

src, dst = sys.argv[1], sys.argv[2]
wb = load_workbook(src)


def cell(v, fmt):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, (int, float)) and "$" in fmt:
        return f"${v:,.2f}"
    return str(v).replace("|", "\\|").replace("\n", "<br>")


out = ["# Sample output: what Paper Chaser produces",
       "",
       "A Markdown copy of [`sample-output.xlsx`](sample-output.xlsx) "
       "([download the Excel file](https://github.com/NormStormm/paper-chaser/raw/main/docs/sample-output.xlsx)), "
       "from one run on the 25 test documents. In Excel, held rows are red, notes yellow, and good insurance green.",
       ""]
for ws in wb:
    out += [f"## {ws.title}", ""]
    rows = list(ws.iter_rows())
    if ws.title == "Summary":
        out += ["| | |", "|---|---|"]
        for r in rows[1:]:
            if r[0].value:
                out.append(f"| {cell(r[0].value, '')} | {cell(r[1].value, r[1].number_format)} |")
    else:
        keep = [i for i, c in enumerate(rows[0]) if c.value not in ("Time", "Tokens in/out", "Seconds", "Reused saved read?", "Model")]
        out.append("| " + " | ".join(str(rows[0][i].value) for i in keep) + " |")
        out.append("|" + "---|" * len(keep))
        for r in rows[1:]:
            out.append("| " + " | ".join(cell(r[i].value, r[i].number_format) for i in keep) + " |")
    out.append("")
out.append("*All vendors, documents and amounts are fictitious test data.*")
open(dst, "w").write("\n".join(out) + "\n")
print("wrote", dst)
